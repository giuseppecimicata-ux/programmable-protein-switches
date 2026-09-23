from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import curve_fit


# ============================================================================
# USER INPUT
# ============================================================================

FILE_PATH = "/path/to/your/input_file.xlsx"

SHEET_NAME = 0
OUTPUT_PATH = "binder_titration_fit.png"

X_LABEL = "Concentration (µM)"
Y_LABEL = "Normalized Activity"
TITLE = "Dose-response curves"

X_LOG = False
Y_LIM = (0, 1.05)

SERIES_NAMES = ["binder_α", "", "Exp3"]
COLORS = ["blue", "orange", "green"]


# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def to_float(value):
    """Convert a value to float, returning NaN when conversion fails."""
    try:
        if value is None:
            return np.nan
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def load_data(path, sheet):
    """
    Load concentration-response data from an Excel worksheet.

    The worksheet is expected to contain:
        column 1: concentration
        column 2: response
        column 3: standard deviation
        column 4: response
        column 5: standard deviation
        ...

    Returns
    -------
    x : numpy.ndarray
        Concentration values.

    experiments : list of tuples
        Each tuple contains (response, standard deviation).
    """
    workbook = load_workbook(
        filename=path,
        data_only=True,
    )

    if isinstance(sheet, int):
        worksheet = workbook.worksheets[sheet]
    else:
        worksheet = workbook[sheet]

    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError("The Excel worksheet is empty.")

    n_columns = max(len(row) for row in rows)

    if n_columns < 3:
        raise ValueError(
            "The worksheet must contain at least three columns: "
            "concentration, response, and standard deviation."
        )

    if (n_columns - 1) % 2 != 0:
        raise ValueError(
            "After the concentration column, the remaining columns must "
            "occur in response / standard-deviation pairs."
        )

    columns = [[] for _ in range(n_columns)]

    for row in rows:
        for index in range(n_columns):
            value = row[index] if index < len(row) else None
            columns[index].append(to_float(value))

    x = np.asarray(columns[0], dtype=float)

    experiments = []

    for index in range(1, n_columns, 2):
        response = np.asarray(columns[index], dtype=float)
        std = np.asarray(columns[index + 1], dtype=float)
        experiments.append((response, std))

    return x, experiments


def hill_increasing(x, ec50, hill_coefficient):
    """Normalized increasing Hill equation."""
    return (
        x ** hill_coefficient
        / (ec50 ** hill_coefficient + x ** hill_coefficient)
    )


def hill_decreasing(x, ec50, hill_coefficient):
    """Normalized decreasing Hill equation."""
    return 1.0 / (
        1.0 + (x / ec50) ** hill_coefficient
    )


def fit_hill_normalized(x, y):
    """
    Fit normalized data using a Hill model with asymptotes fixed at 0 and 1.

    Returns
    -------
    model : callable
        Selected Hill model.

    parameters : numpy.ndarray
        Fitted parameters [EC50, Hill coefficient].

    ec50 : float
        Apparent EC50.

    hill_coefficient : float
        Fitted Hill coefficient.
    """
    if np.any(x <= 0):
        raise ValueError(
            "All concentrations must be greater than zero for Hill fitting."
        )

    increasing = y[-1] > y[0]

    ec50_guess = np.median(x)
    hill_guess = 1.0

    bounds = (
        [np.min(x) * 1e-6, 0.01],
        [np.max(x) * 1e6, 10.0],
    )

    if increasing:
        model = hill_increasing
    else:
        model = hill_decreasing

    parameters, _ = curve_fit(
        model,
        x,
        y,
        p0=[ec50_guess, hill_guess],
        bounds=bounds,
        maxfev=20000,
    )

    ec50 = parameters[0]
    hill_coefficient = parameters[1]

    return model, parameters, ec50, hill_coefficient


# ============================================================================
# MAIN ANALYSIS
# ============================================================================

def main():

    file_path = Path(FILE_PATH)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Input file not found: {file_path}"
        )

    x, experiments = load_data(
        file_path,
        SHEET_NAME,
    )

    # Remove rows containing invalid concentrations.
    valid = np.isfinite(x)

    for response, std in experiments:
        valid &= np.isfinite(response) & np.isfinite(std)

    x = x[valid]
    experiments = [
        (response[valid], std[valid])
        for response, std in experiments
    ]

    if len(x) == 0:
        raise ValueError("No valid data points were found.")

    # Sort by concentration.
    order = np.argsort(x)

    x = x[order]
    experiments = [
        (response[order], std[order])
        for response, std in experiments
    ]

    fig, ax = plt.subplots(figsize=(7, 5))

    for index, (response, std) in enumerate(experiments):

        name = (
            SERIES_NAMES[index]
            if index < len(SERIES_NAMES)
            else f"Series {index + 1}"
        )

        color = (
            COLORS[index]
            if index < len(COLORS)
            else None
        )

        # Plot experimental data and standard deviation.
        ax.errorbar(
            x,
            response,
            yerr=std,
            fmt="o",
            linestyle="none",
            label=name,
            color=color,
            capsize=3,
            markersize=5,
            zorder=3,
        )

        try:
            model, parameters, ec50, hill_coefficient = (
                fit_hill_normalized(x, response)
            )

            x_fit = np.linspace(
                np.min(x),
                np.max(x),
                400,
            )

            y_fit = model(
                x_fit,
                *parameters,
            )

            ax.plot(
                x_fit,
                y_fit,
                color=color,
                linewidth=2.2,
                label=f"{name} fit (Kd,app={ec50:.3g})",
                zorder=2,
            )

            print(
                f"{name}: Kd,app = {ec50:.6g}, "
                f"Hill coefficient = {hill_coefficient:.3g}"
            )

        except Exception as error:
            print(
                f"Hill fit failed for {name}: {error}"
            )

    if X_LOG:
        ax.set_xscale("log")

    ax.set_xlabel(X_LABEL)
    ax.set_ylabel(Y_LABEL)

    if TITLE:
        ax.set_title(TITLE)

    if Y_LIM is not None:
        ax.set_ylim(Y_LIM)

    ax.set_xlim(left=0)

    ax.legend()
    ax.grid(False)

    output_path = Path(OUTPUT_PATH)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    fig.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )

    plt.show()

    print(f"Figure saved to: {output_path}")


if __name__ == "__main__":
    main()
