from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import curve_fit


# ============================================================================
# USER INPUT
# ============================================================================

FILE_PATH = "/path/to/your/input_file.xlsx"


SHEET_NAME = 0
OUTPUT_PATH = "countertitration_fit.png"

X_LABEL = "Switch-release concentration (µM)"
Y_LABEL = "Normalized Activity"
TITLE = "Countertitration"

X_LOG = True
Y_LIM = (0, 1.05)

# Color used for both experimental data and fitted curve
PLOT_COLOR = "tab:pink"

# Total binder concentration used in the countertitration experiment (µM)
BINDER_FIXED_UM = 150.0

# Apparent Kd obtained from the independent binder titration (µM)
ACTIVATION_KD_APP_UM = 18.1

# Wide bounds for the apparent competition Kd
KD_MIN = 1e-12
KD_MAX = 1e12


# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def to_float(value):
    """Convert a value to float, returning NaN when conversion fails."""
    try:
        if value is None:
            return np.nan
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def load_data(path, sheet):
    """
    Load countertitration data from an Excel worksheet.

    Expected format:
        column 1: switch-release concentration
        column 2: normalized response
    """
    workbook = load_workbook(
        filename=path,
        data_only=True,
    )

    if isinstance(sheet, int):
        worksheet = workbook.worksheets[sheet]
    else:
        worksheet = workbook[sheet]

    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError("The Excel worksheet is empty.")

    n_columns = max(len(row) for row in rows)

    if n_columns < 2:
        raise ValueError(
            "The worksheet must contain at least two columns: "
            "concentration and response."
        )

    columns = [[] for _ in range(n_columns)]

    for row in rows:
        for index in range(n_columns):
            value = row[index] if index < len(row) else None
            columns[index].append(to_float(value))

    x = np.asarray(columns[0], dtype=float)
    y = np.asarray(columns[1], dtype=float)

    valid = np.isfinite(x) & np.isfinite(y)

    x = x[valid]
    y = y[valid]

    order = np.argsort(x)

    return x[order], y[order]


def quadratic_complex(a_total, b_total, kd):
    """
    Calculate the concentration of a 1:1 complex from total concentrations.

    Equilibrium:
        A + B <-> AB
    """
    a_total = np.asarray(a_total, dtype=float)
    b_total = np.asarray(b_total, dtype=float)

    discriminant = (
        (a_total + b_total + kd) ** 2
        - 4.0 * a_total * b_total
    )

    # Protect against very small negative values caused by
    # floating-point numerical precision.
    discriminant = np.maximum(discriminant, 0.0)

    return 0.5 * (
        (a_total + b_total + kd)
        - np.sqrt(discriminant)
    )


def free_binder_after_sequestration(
    binder_total,
    switch_release_total,
    kd_competition,
):
    """
    Calculate free binder after sequestration by switch-release.

    Equilibrium:
        Binder + Switch-release <-> Complex
    """
    complex_concentration = quadratic_complex(
        binder_total,
        switch_release_total,
        kd_competition,
    )

    return np.maximum(
        binder_total - complex_concentration,
        0.0,
    )


def response_from_free_binder(
    binder_free,
    bottom,
    top,
    hill_coefficient,
    increasing_response=True,
):
    """
    Calculate the normalized response from free binder concentration.

    The apparent Kd obtained from the independent binder titration
    is used as a fixed parameter.
    """
    closed_fraction = (
        binder_free ** hill_coefficient
        / (
            ACTIVATION_KD_APP_UM ** hill_coefficient
            + binder_free ** hill_coefficient
        )
    )

    if increasing_response:
        return bottom + (top - bottom) * (1.0 - closed_fraction)

    return bottom + (top - bottom) * closed_fraction


def competition_model(
    switch_release,
    kd_competition,
    bottom,
    top,
    hill_coefficient,
    increasing_response=True,
):
    """
    Calculate the countertitration response.

    Free binder is first calculated after sequestration by
    switch-release and then converted into the normalized response.
    """
    binder_free = free_binder_after_sequestration(
        BINDER_FIXED_UM,
        switch_release,
        kd_competition,
    )

    return response_from_free_binder(
        binder_free,
        bottom,
        top,
        hill_coefficient,
        increasing_response=increasing_response,
    )


# ============================================================================
# MAIN ANALYSIS
# ============================================================================

def main():

    file_path = Path(FILE_PATH)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Input file not found: {file_path}"
        )

    x, y = load_data(
        file_path,
        SHEET_NAME,
    )

    if len(x) == 0:
        raise ValueError("No valid data points were found.")

    if X_LOG and np.any(x <= 0):
        raise ValueError(
            "All switch-release concentrations must be greater than zero "
            "when X_LOG=True."
        )

    # Automatically determine the direction of the response.
    increasing_response = y[-1] > y[0]

    def model(x_data, kd, bottom, top, h):
        return competition_model(
            x_data,
            kd,
            bottom,
            top,
            h,
            increasing_response=increasing_response,
        )

    # Initial parameter estimates.
    initial_parameters = [
        np.median(x),   # apparent competition Kd
        np.min(y),      # bottom
        np.max(y),      # top
        1.0,            # Hill-like exponent
    ]

    parameter_bounds = (
        [KD_MIN, -np.inf, -np.inf, 0.05],
        [KD_MAX, np.inf, np.inf, 10.0],
    )

    fitted_parameters, covariance = curve_fit(
        model,
        x,
        y,
        p0=initial_parameters,
        bounds=parameter_bounds,
        maxfev=50000,
    )

    kd_competition, bottom, top, hill_coefficient = fitted_parameters

    parameter_errors = np.sqrt(
        np.diag(covariance)
    )

    if X_LOG:
        x_fit = np.logspace(
            np.log10(np.min(x)),
            np.log10(np.max(x)),
            400,
        )
    else:
        x_fit = np.linspace(
            np.min(x),
            np.max(x),
            400,
        )

    y_fit = model(
        x_fit,
        *fitted_parameters,
    )


    # =========================================================================
    # REPORT FIT RESULTS
    # =========================================================================

    print("\n=== COUNTERTITRATION FIT ===")

    print(
        f"Fixed binder concentration = "
        f"{BINDER_FIXED_UM:.3g} µM"
    )

    print(
        f"Fixed activation Kd,app = "
        f"{ACTIVATION_KD_APP_UM:.3g} µM"
    )

    print(
        f"Competition Kd,app = "
        f"{kd_competition:.6g} ± "
        f"{parameter_errors[0]:.2g} µM"
    )

    print(
        f"Bottom = {bottom:.6g} ± "
        f"{parameter_errors[1]:.2g}"
    )

    print(
        f"Top = {top:.6g} ± "
        f"{parameter_errors[2]:.2g}"
    )

    print(
        f"Hill-like exponent = "
        f"{hill_coefficient:.6g} ± "
        f"{parameter_errors[3]:.2g}"
    )


    # =========================================================================
    # PLOT
    # =========================================================================

    fig, ax = plt.subplots(figsize=(7, 5))

    # Experimental data
    ax.plot(
        x,
        y,
        "o",
        color=PLOT_COLOR,
        markersize=5,
        label="Data",
        zorder=3,
    )

    # Fitted curve
    ax.plot(
        x_fit,
        y_fit,
        color=PLOT_COLOR,
        linewidth=2.2,
        label=f"Fit (Kd,app={kd_competition:.3g} µM)",
        zorder=2,
    )

    if X_LOG:
        ax.set_xscale("log")

    ax.set_xlabel(X_LABEL)
    ax.set_ylabel(Y_LABEL)

    if TITLE:
        ax.set_title(TITLE)

    if Y_LIM is not None:
        ax.set_ylim(Y_LIM)

    ax.set_xlim(
        left=np.min(x) * 0.9,
        right=np.max(x) * 1.1,
    )

    ax.legend()
    ax.grid(False)

    output_path = Path(OUTPUT_PATH)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    fig.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )

    plt.show()

    print(f"Figure saved to: {output_path}")


if __name__ == "__main__":
    main() 
